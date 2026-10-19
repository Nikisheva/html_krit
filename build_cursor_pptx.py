#!/usr/bin/env python3
"""Training deck from the KRIT back-office Cursor instruction."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

W = Inches(13.333)
H = Inches(7.5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x1E, 0x1E)
MUTED = RGBColor(0x77, 0x77, 0x77)
ACCENT = RGBColor(0x38, 0x9B, 0xFF)
LINE = RGBColor(0xE0, 0xE3, 0xE8)
SOFT = RGBColor(0xF4, 0xF8, 0xFC)
HEAD = "Uncage"
BODY = "Arial"
FOOTER = "Бэк-офис и ОЦО · AI Challenge"


def set_run(run, text, size, color, font=BODY, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    run.font.bold = bold
    rPr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", font)


def add_box(slide, l, t, w, h, fill, line=None, line_w=Pt(0.75)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = line_w
    return sh


def add_bar(slide, l, t, w, fill=ACCENT, h=Pt(3)):
    return add_box(slide, l, t, w, h, fill)


def add_text(slide, l, t, w, h, text, size, color, font=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size, color, font=font, bold=bold)
    return sh


def add_lines(slide, l, t, w, h, lines, size, color, font=BODY, bold=False, space_after=8, anchor=MSO_ANCHOR.TOP):
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        run = p.add_run()
        set_run(run, line, size, color, font=font, bold=bold)
    return sh


def card(slide, l, t, w, h, fill=WHITE):
    add_box(slide, l, t, w, h, fill, LINE)
    add_bar(slide, l, t, w)


def header(slide, kicker, title, num):
    add_box(slide, Inches(0.28), Inches(1.18), Pt(2.5), Inches(5.55), ACCENT)
    add_text(slide, Inches(0.55), Inches(0.22), Inches(10.4), Inches(0.32), kicker, 12, ACCENT)
    add_text(slide, Inches(0.55), Inches(0.48), Inches(10.6), Inches(0.55), title, 26, INK, font=HEAD)
    slide.shapes.add_picture("assets/pptx/logo-grey.png", Inches(11.85), Inches(0.28), height=Inches(0.42))
    add_text(slide, Inches(0.55), Inches(7.12), Inches(9), Inches(0.28), FOOTER, 11, ACCENT)
    add_text(slide, Inches(12.15), Inches(7.12), Inches(0.7), Inches(0.28), num, 11, MUTED, align=PP_ALIGN.RIGHT)


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, WHITE)
    return s


def blue_panel(slide):
    builder = slide.shapes.build_freeform(Inches(10.55), Inches(0))
    builder.add_line_segments([(W, Inches(0)), (W, H), (Inches(8.82), H)], close=True)
    geo = builder.convert_to_shape()
    geo.shadow.inherit = False
    geo.fill.solid()
    geo.fill.fore_color.rgb = ACCENT
    geo.line.fill.background()
    slide.shapes.add_picture("assets/pptx/logo-white.png", Inches(11.85), Inches(0.32), height=Inches(0.42))


def cover(prs):
    s = blank(prs)
    blue_panel(s)
    add_box(s, Inches(0.7), Inches(1.35), Inches(5.35), Inches(0.36), SOFT)
    add_text(s, Inches(0.82), Inches(1.37), Inches(5.1), Inches(0.32), "Обучение · бэк-офис и ОЦО", 13, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(0.7), Inches(1.95), Inches(8.2), Inches(1.5), "Работа\nс Cursor", 40, INK, font=HEAD)
    add_text(
        s,
        Inches(0.7),
        Inches(4.15),
        Inches(7.5),
        Inches(1.35),
        "Помощник на вашем компьютере: разбирает задачу, готовит правки и помогает проверить результат. За то, что попросили, согласовали и проверили, отвечаете вы.",
        18,
        INK,
    )
    add_text(s, Inches(0.7), Inches(6.85), Inches(7), Inches(0.3), "Откройте инструкцию перед началом задачи", 14, MUTED)


def what(prs):
    s = blank(prs)
    header(s, "С чего начать", "Помощник внутри вашей папки", "02")
    add_text(
        s,
        Inches(0.55),
        Inches(1.2),
        Inches(12.2),
        Inches(0.55),
        "Описываете задачу обычными словами. Cursor помогает выполнить её в открытом проекте на вашем компьютере.",
        16,
        INK,
    )
    items = [
        ("01", "Найти", "Инструкцию или описание процесса в проекте."),
        ("02", "Подготовить", "Правку во внутренней системе или автоматизации."),
        ("03", "Показать", "Что именно изменилось после работы."),
        ("04", "Сохранить", "Понятную историю этих изменений."),
    ]
    cw, ch, gap = Inches(2.95), Inches(2.35), Inches(0.16)
    y = Inches(1.9)
    for i, (num, title, body) in enumerate(items):
        x = Inches(0.55) + i * (cw + gap)
        card(s, x, y, cw, ch)
        add_text(s, x + Inches(0.18), y + Inches(0.22), Inches(2.5), Inches(0.32), num, 14, ACCENT, font=HEAD)
        add_text(s, x + Inches(0.18), y + Inches(0.58), Inches(2.55), Inches(0.4), title, 20, INK, font=HEAD)
        add_text(s, x + Inches(0.18), y + Inches(1.15), Inches(2.58), Inches(0.95), body, 14, MUTED)
    add_box(s, Inches(0.55), Inches(4.5), Inches(12.22), Inches(2.15), SOFT)
    add_text(s, Inches(0.8), Inches(4.68), Inches(11.7), Inches(0.35), "Работаем локально, не в Cloud", 18, INK, font=HEAD)
    add_text(
        s,
        Inches(0.8),
        Inches(5.15),
        Inches(11.7),
        Inches(1.25),
        "Файлы лежат на вашем компьютере. Cloud для рабочих проектов и корпоративной информации не используем: на одной учётной записи могут работать несколько человек. Не уверены, где вы сейчас, — остановитесь и спросите IT.",
        16,
        INK,
    )


def plan(prs):
    s = blank(prs)
    header(s, "Как начинать задачу", "Сначала план, потом изменения", "03")
    card(s, Inches(0.55), Inches(1.25), Inches(7.55), Inches(4.15))
    add_lines(
        s,
        Inches(0.8),
        Inches(1.5),
        Inches(7.05),
        Inches(3.7),
        [
            "01   Опишите задачу своими словами.",
            "02   Попросите изучить, как сделано сейчас, и пока не менять файлы.",
            "03   Попросите пошаговый план.",
            "04   Сами прочитайте план: всё ли понятно и нет ли лишнего.",
            "05   Только после этого разрешите изменения.",
        ],
        16,
        INK,
        space_after=14,
    )
    card(s, Inches(8.28), Inches(1.25), Inches(4.5), Inches(4.15), SOFT)
    add_text(s, Inches(8.5), Inches(1.48), Inches(4.1), Inches(0.7), "Задача сложная, если", 18, INK, font=HEAD)
    add_lines(
        s,
        Inches(8.5),
        Inches(2.25),
        Inches(4.05),
        Inches(2.9),
        [
            "Затронуто несколько файлов, систем или отделов.",
            "Можно сломать уже работающий процесс.",
            "Нужно менять доступы, данные или настройки.",
            "Вы сами не до конца понимаете, что изменится.",
        ],
        14,
        MUTED,
        space_after=10,
    )
    add_text(
        s,
        Inches(0.55),
        Inches(5.6),
        Inches(12.2),
        Inches(1.15),
        "Формула запроса: контекст → что сделать → где можно менять → что нельзя → как проверить. Простую задачу можно пройти короче. Сложную разбирайте самым сильным помощником.",
        16,
        INK,
    )


def bounds(prs):
    s = blank(prs)
    header(s, "Границы", "Что помощнику не отдаём", "04")
    items = [
        (
            "Секреты",
            "Пароли, коды и ключи не пишем в чат, не вставляем в запрос и не кладём в историю проекта. Даже если попросите показать пароль. Нужен доступ — Cursor готовит пустой локальный файл, значения вносите вы.",
        ),
        (
            "Рабочая система",
            "DEV — здесь пробуем. TEST — здесь проверяем. PROD — система, в которой идёт реальная работа. В неё без отдельного процесса не лезем. Неясно, где вы, — остановитесь и спросите IT.",
        ),
        (
            "Ядро",
            "Заводскую часть системы, чужие библиотеки и файлы не по вашей задаче не меняем. Если без этого никак, Cursor останавливается, объясняет риск и ждёт вашего решения или решения IT.",
        ),
    ]
    cw, ch, gap = Inches(3.95), Inches(4.15), Inches(0.18)
    y = Inches(1.28)
    for i, (title, body) in enumerate(items):
        x = Inches(0.55) + i * (cw + gap)
        card(s, x, y, cw, ch)
        add_text(s, x + Inches(0.24), y + Inches(0.28), Inches(3.45), Inches(0.7), title, 22, INK, font=HEAD)
        add_text(s, x + Inches(0.24), y + Inches(1.15), Inches(3.45), Inches(2.7), body, 15, MUTED)
    add_text(
        s,
        Inches(0.55),
        Inches(5.65),
        Inches(12.2),
        Inches(1.1),
        "Тоже стоп: много удалений, смена доступов, база данных, массовое изменение данных. Если остаётся вопрос «а можно ли это менять?» — менять ещё нельзя.",
        16,
        INK,
    )


def close(prs):
    s = blank(prs)
    header(s, "Как заканчиваем", "«Готово» ещё не значит готово", "05")
    card(s, Inches(0.55), Inches(1.25), Inches(6.0), Inches(4.35))
    add_text(s, Inches(0.8), Inches(1.45), Inches(5.5), Inches(0.4), "Перед началом", 18, INK, font=HEAD)
    add_lines(
        s,
        Inches(0.8),
        Inches(2.0),
        Inches(5.5),
        Inches(3.3),
        [
            "Работаю локально, не в Cloud.",
            "Понимаю: это DEV, TEST или PROD.",
            "История проекта подключена. Её настраивает IT.",
            "Правило про пароли и ключи включено.",
            "Для сложной задачи сначала есть план.",
            "Понимаю, что именно собираюсь менять.",
        ],
        15,
        INK,
        space_after=8,
    )
    card(s, Inches(6.75), Inches(1.25), Inches(6.02), Inches(4.35), SOFT)
    add_text(s, Inches(7.0), Inches(1.45), Inches(5.5), Inches(0.4), "После изменений", 18, INK, font=HEAD)
    add_lines(
        s,
        Inches(7.0),
        Inches(2.0),
        Inches(5.5),
        Inches(3.3),
        [
            "Посмотрел, какие файлы изменились.",
            "Секреты в историю не попали.",
            "Основная функция работает.",
            "Соседние процессы не сломались.",
            "Ядро и чужие файлы не тронуты.",
            "Версия сохранена и отправлена в общий проект.",
        ],
        15,
        INK,
        space_after=8,
    )
    add_text(
        s,
        Inches(0.55),
        Inches(5.8),
        Inches(12.2),
        Inches(1.0),
        "Сообщение «готово» не заменяет проверку. Список изменений, сохранение версии и отправка в общий проект — это diff, commit и push. Сомневаетесь — зовём IT.",
        16,
        INK,
    )


def quote(prs):
    s = blank(prs)
    blue_panel(s)
    add_text(s, Inches(0.7), Inches(1.15), Inches(8), Inches(0.35), "Вместо заключения", 14, ACCENT)
    add_text(
        s,
        Inches(0.7),
        Inches(1.7),
        Inches(8.3),
        Inches(2.6),
        "«Будущее уже здесь.\nПросто оно распределено\nнеравномерно.»",
        32,
        INK,
        font=HEAD,
    )
    add_text(s, Inches(0.7), Inches(4.55), Inches(7.5), Inches(0.35), "Уильям Гибсон", 16, INK)
    add_text(s, Inches(0.7), Inches(4.95), Inches(7.8), Inches(0.3), "The Economist, 4 декабря 2003. Перевод.", 13, MUTED)
    add_text(
        s,
        Inches(0.7),
        Inches(5.7),
        Inches(7.8),
        Inches(0.9),
        "Помощник может смениться. Проект остаётся на вашем компьютере, а ответственность за результат — на вас.",
        16,
        INK,
    )


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Работа с Cursor"
    prs.core_properties.subject = "Обучение для бэк-офиса и ОЦО"
    cover(prs)
    what(prs)
    plan(prs)
    bounds(prs)
    close(prs)
    quote(prs)
    out = "Cursor для AI Challenge.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
